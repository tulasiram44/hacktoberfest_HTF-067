import sys
from docx import Document
from pptx import Presentation

def extract_docx(file_path):
    doc = Document(file_path)
    text = []
    for para in doc.paragraphs:
        if para.text.strip():
            text.append(para.text)
    return '\n'.join(text)

def extract_pptx(file_path):
    prs = Presentation(file_path)
    text = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text.append(shape.text)
    return '\n'.join(text)

if __name__ == "__main__":
    import os
    with open("extracted_data.txt", "w", encoding="utf-8") as f:
        f.write("=== DOCX ===\n")
        docx_path = r"C:\Users\Neha\OneDrive\Desktop\SIH\Sahaya_Technical_Reference.docx"
        if os.path.exists(docx_path):
            f.write(extract_docx(docx_path) + "\n")
        else:
            f.write("Docx not found\n")
            
        f.write("\n=== PPTX ===\n")
        pptx_path = r"C:\Users\Neha\OneDrive\Desktop\SIH26-A0H-T133-SIH26089_Presentation_imp.pptx"
        if os.path.exists(pptx_path):
            f.write(extract_pptx(pptx_path) + "\n")
        else:
            f.write("Pptx not found\n")
